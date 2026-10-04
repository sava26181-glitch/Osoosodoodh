import asyncio
import csv
import io
import json
import logging
import os
import re
import zipfile
from pathlib import Path

from aiohttp import web
from aiogram import Bot, Dispatcher, F, Router
from aiogram.filters import Command, CommandStart
from aiogram.types import BufferedInputFile, Message
from aiogram.webhook.aiohttp_server import SimpleRequestHandler, setup_application
from anthropic import AsyncAnthropic

logging.basicConfig(level=logging.INFO)
log = logging.getLogger(__name__)

BOT_TOKEN = os.environ["TELEGRAM_BOT_TOKEN"]
ANTHROPIC_API_KEY = os.environ["ANTHROPIC_API_KEY"]
WEBHOOK_SECRET = os.environ["WEBHOOK_SECRET"]
BASE_WEBHOOK_URL = os.getenv("BASE_WEBHOOK_URL") or os.getenv("RENDER_EXTERNAL_URL")

if not BASE_WEBHOOK_URL:
    raise RuntimeError("BASE_WEBHOOK_URL or RENDER_EXTERNAL_URL is required")

WEBHOOK_PATH = os.getenv("WEBHOOK_PATH", "/telegram/webhook")
PORT = int(os.getenv("PORT", "10000"))
CLAUDE_MODEL = os.getenv("CLAUDE_MODEL", "claude-sonnet-4-5")
MAX_TOKENS = int(os.getenv("MAX_TOKENS", "12000"))
SYSTEM_PROMPT = os.getenv(
    "SYSTEM_PROMPT",
    """Ты Claude, работающий внутри Telegram-бота.
Отвечай на русском, если пользователь не просит другой язык.
Если пользователь просит создать файл, верни результат строго в формате FILE_REQUEST.
Не добавляй markdown fences вокруг содержимого файла.
"""
)

bot = Bot(BOT_TOKEN)
dp = Dispatcher()
router = Router()
dp.include_router(router)
claude = AsyncAnthropic(api_key=ANTHROPIC_API_KEY)

histories: dict[int, list[dict[str, str]]] = {}
MAX_HISTORY = 20

ALLOWED_EXTENSIONS = {
    ".txt", ".md", ".py", ".js", ".ts", ".html", ".css", ".json",
    ".csv", ".xml", ".sql", ".yaml", ".yml", ".sh", ".bat", ".java",
    ".cpp", ".c", ".h", ".php", ".go", ".rs", ".jsx", ".tsx"
}

FILE_RE = re.compile(
    r"FILE_REQUEST\s*\n"
    r"filename\s*:\s*(?P<filename>[^\n]+)\n"
    r"content\s*:\s*\n(?P<content>[\s\S]*)",
    re.IGNORECASE
)


def get_history(chat_id: int):
    return histories.setdefault(chat_id, [])


def add_history(chat_id: int, role: str, content: str):
    h = get_history(chat_id)
    h.append({"role": role, "content": content})
    if len(h) > MAX_HISTORY:
        del h[:-MAX_HISTORY]


async def ask_claude(chat_id: int, prompt: str) -> str:
    add_history(chat_id, "user", prompt)

    response = await claude.messages.create(
        model=CLAUDE_MODEL,
        max_tokens=MAX_TOKENS,
        system=SYSTEM_PROMPT,
        messages=get_history(chat_id),
    )

    text = "\n".join(
        block.text
        for block in response.content
        if getattr(block, "type", None) == "text"
    ).strip()

    add_history(chat_id, "assistant", text)
    return text


def clean_filename(name: str, fallback="file.txt") -> str:
    name = name.strip().replace("\\", "_").replace("/", "_")
    name = re.sub(r"[\x00-\x1f<>:\"|?*]", "_", name)
    name = name[:120]
    return name or fallback


def make_document(filename: str, data: bytes) -> BufferedInputFile:
    return BufferedInputFile(data, filename=filename)


def text_bytes(text: str) -> bytes:
    return text.encode("utf-8")


def make_csv_from_json(text: str) -> bytes:
    data = json.loads(text)
    if not isinstance(data, list):
        raise ValueError("Для CSV нужен JSON-массив объектов.")

    if not data:
        return b""

    keys = []
    for item in data:
        if isinstance(item, dict):
            for key in item:
                if key not in keys:
                    keys.append(key)

    out = io.StringIO()
    writer = csv.DictWriter(out, fieldnames=keys)
    writer.writeheader()
    for item in data:
        writer.writerow(item if isinstance(item, dict) else {})
    return out.getvalue().encode("utf-8-sig")


def make_xlsx(text: str) -> bytes:
    from openpyxl import Workbook

    data = json.loads(text)
    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet1"

    if isinstance(data, list) and data and all(isinstance(x, dict) for x in data):
        keys = []
        for item in data:
            for key in item:
                if key not in keys:
                    keys.append(key)
        ws.append(keys)
        for item in data:
            ws.append([item.get(k, "") for k in keys])
    elif isinstance(data, list):
        for row in data:
            ws.append(row if isinstance(row, list) else [row])
    elif isinstance(data, dict):
        ws.append(["key", "value"])
        for k, v in data.items():
            ws.append([k, v])
    else:
        ws.append(["value"])
        ws.append([data])

    output = io.BytesIO()
    wb.save(output)
    return output.getvalue()


def make_docx(text: str) -> bytes:
    from docx import Document

    doc = Document()
    for paragraph in text.split("\n\n"):
        doc.add_paragraph(paragraph)
    output = io.BytesIO()
    doc.save(output)
    return output.getvalue()


def make_pdf(text: str) -> bytes:
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    from xml.sax.saxutils import escape

    output = io.BytesIO()

    # DejaVu поддерживает кириллицу и обычно доступен в Render.
    font_paths = [
        "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf",
        "/usr/share/fonts/dejavu/DejaVuSans.ttf",
    ]
    font_path = next((p for p in font_paths if Path(p).exists()), None)

    if font_path:
        pdfmetrics.registerFont(TTFont("DejaVu", font_path))
        styles = getSampleStyleSheet()
        styles["Normal"].fontName = "DejaVu"
        styles["Normal"].fontSize = 10
        styles["Normal"].leading = 14
    else:
        styles = getSampleStyleSheet()

    doc = SimpleDocTemplate(
        output,
        pagesize=A4,
        rightMargin=45,
        leftMargin=45,
        topMargin=45,
        bottomMargin=45,
    )

    story = []
    for paragraph in text.split("\n"):
        value = escape(paragraph) if paragraph else " "
        story.append(Paragraph(value, styles["Normal"]))
        story.append(Spacer(1, 5))

    doc.build(story)
    return output.getvalue()


def make_zip(filename: str, files: list[tuple[str, bytes]]) -> bytes:
    output = io.BytesIO()
    with zipfile.ZipFile(output, "w", zipfile.ZIP_DEFLATED) as zf:
        for name, data in files:
            zf.writestr(name, data)
    return output.getvalue()


def parse_file_request(answer: str):
    match = FILE_RE.search(answer)
    if not match:
        return None

    filename = clean_filename(match.group("filename"))
    content = match.group("content").strip()

    # Убираем случайные markdown fences.
    if content.startswith("```") and content.endswith("```"):
        lines = content.splitlines()
        if len(lines) >= 2:
            content = "\n".join(lines[1:-1])

    return filename, content


async def create_requested_file(answer: str):
    parsed = parse_file_request(answer)
    if not parsed:
        return None

    filename, content = parsed
    suffix = Path(filename).suffix.lower()

    if suffix == ".xlsx":
        data = make_xlsx(content)
    elif suffix == ".docx":
        data = make_docx(content)
    elif suffix == ".pdf":
        data = make_pdf(content)
    elif suffix == ".csv":
        # Claude can return JSON for structured CSV generation.
        try:
            data = make_csv_from_json(content)
        except Exception:
            data = content.encode("utf-8-sig")
    elif suffix == ".zip":
        # For ZIP, content can be a JSON object:
        # {"file.txt":"...", "main.py":"..."}
        try:
            files = json.loads(content)
            if not isinstance(files, dict):
                raise ValueError
            items = [(clean_filename(k), str(v).encode("utf-8")) for k, v in files.items()]
            data = make_zip(filename, items)
        except Exception:
            data = make_zip(filename, [("content.txt", content.encode("utf-8"))])
    else:
        if suffix and suffix not in ALLOWED_EXTENSIONS:
            filename = f"{Path(filename).stem}.txt"
        data = text_bytes(content)

    return filename, data


@router.message(CommandStart())
async def start(message: Message):
    await message.answer(
        "Привет! Я Claude-файловый агент.\n\n"
        "Могу отвечать на вопросы и создавать настоящие файлы:\n"
        "TXT, MD, PY, JS, HTML, CSS, JSON, CSV, XLSX, DOCX, PDF, ZIP и другие.\n\n"
        "Пример:\n"
        "«Создай Excel с расходами за месяц и отправь мне»\n\n"
        "/reset — очистить память"
    )


@router.message(Command("reset"))
async def reset(message: Message):
    histories.pop(message.chat.id, None)
    await message.answer("Память диалога очищена.")


@router.message(Command("help"))
async def help_command(message: Message):
    await message.answer(
        "Просто опиши, что нужно создать.\n\n"
        "Например:\n"
        "• «Создай Python-скрипт для переименования файлов»\n"
        "• «Сделай Excel-таблицу расходов»\n"
        "• «Создай PDF-инструкцию на русском»\n"
        "• «Сделай Word-документ с резюме»\n"
        "• «Создай сайт и упакуй HTML/CSS/JS в ZIP»"
    )


@router.message(F.document)
async def incoming_document(message: Message):
    await message.answer(
        "Входящие файлы пока не анализирую. "
        "Могу создавать и отправлять файлы по твоему описанию."
    )


@router.message(F.text)
async def text_handler(message: Message):
    prompt = message.text.strip()
    if not prompt:
        return

    await message.bot.send_chat_action(message.chat.id, "typing")

    try:
        answer = await ask_claude(message.chat.id, prompt)
        generated = await create_requested_file(answer)
    except Exception:
        log.exception("Generation error")
        await message.answer(
            "Произошла ошибка при генерации. Проверь API-ключи и логи Render."
        )
        return

    if generated:
        filename, data = generated
        if len(data) > 49 * 1024 * 1024:
            await message.answer(
                "Файл получился слишком большим для отправки через Telegram."
            )
            return

        await message.answer_document(
            make_document(filename, data),
            caption=f"Готово: {filename}"
        )
        return

    # Если это обычный ответ — отправляем текст.
    for i in range(0, len(answer), 4000):
        await message.answer(answer[i:i + 4000])


async def health(request: web.Request):
    return web.json_response({"status": "ok"})


async def on_startup(bot: Bot):
    webhook_url = f"{BASE_WEBHOOK_URL.rstrip('/')}{WEBHOOK_PATH}"
    await bot.set_webhook(
        webhook_url,
        secret_token=WEBHOOK_SECRET,
        drop_pending_updates=True,
        allowed_updates=dp.resolve_used_update_types(),
    )
    log.info("Webhook: %s", webhook_url)


async def on_shutdown(bot: Bot):
    try:
        await bot.delete_webhook(drop_pending_updates=False)
    finally:
        await claude.close()
        await bot.session.close()


def create_app():
    app = web.Application()
    app.router.add_get("/", health)
    app.router.add_get("/health", health)

    handler = SimpleRequestHandler(
        dispatcher=dp,
        bot=bot,
        secret_token=WEBHOOK_SECRET,
        handle_in_background=True,
    )
    handler.register(app, path=WEBHOOK_PATH)

    dp.startup.register(on_startup)
    dp.shutdown.register(on_shutdown)
    setup_application(app, dp, bot=bot)

    return app


if __name__ == "__main__":
    web.run_app(create_app(), host="0.0.0.0", port=PORT)
